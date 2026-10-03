#!/usr/bin/env python3
"""
Export complet de fin de course — LECTURE SEULE.
Classeur Excel (si openpyxl dispo, sinon CSV) avec : Résultats (2 décomptes),
Vélos, Écuries, Historique points, Achats, Corrections, Passages.

Deux décomptes de tours :
 - « App »     : ce que le système a compté pendant la course.
 - « Règle »   : recompté hors ligne avec un écart mini de MIN_GAP s entre 2
                 passages d'un même vélo et un arrêt à FIN (heure de Bruxelles).
Bonus/malus/corrections sont ajoutés aux deux.
"""
import csv, io, json, subprocess, sys, urllib.request
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

MIN_GAP = int(sys.argv[1]) if len(sys.argv) > 1 else 120
FIN = sys.argv[2] if len(sys.argv) > 2 else '2026-10-03 16:30:00'
TZ = ZoneInfo('Europe/Brussels')
fin_utc = datetime.fromisoformat(FIN).replace(tzinfo=TZ).astimezone(ZoneInfo('UTC')).replace(tzinfo=None)

def sql(q):
    out = subprocess.run(['docker', 'compose', 'exec', '-T', 'db', 'psql', '-U', 'ft', '-d', 'ft', '-c', f'COPY ({q}) TO STDOUT WITH CSV HEADER'],
                         cwd='/opt/ft', capture_output=True, text=True, check=True).stdout
    return list(csv.DictReader(io.StringIO(out)))

def local(ts):
    if not ts: return ''
    return datetime.fromisoformat(ts[:19]).replace(tzinfo=ZoneInfo('UTC')).astimezone(TZ).strftime('%H:%M:%S')

live = json.load(urllib.request.urlopen('http://127.0.0.1:8092/api/live', timeout=15))
teams = {t['id']: t for t in live['teams']}
groups = {c['key']: c['name'] for c in live['contests']}
passings = sql('SELECT "dossardNumber", timestamp, "createdAt", source, array_to_string(mats, \' + \') AS mats FROM "RaceLapEvent" ORDER BY "dossardNumber", timestamp')

# --- Recomptage "règle" -------------------------------------------------------
by_bike = {}
for p in passings:
    by_bike.setdefault(int(p['dossardNumber']), []).append(datetime.fromisoformat(p['timestamp'][:23]))
rule = {}
merged_detail = {}
for n, times in by_bike.items():
    last, count, merged, late = None, 0, 0, 0
    for t in times:
        if t > fin_utc: late += 1; continue
        if last is not None and (t - last).total_seconds() < MIN_GAP: merged += 1; continue
        count += 1; last = t
    rule[n] = count; merged_detail[n] = (merged, late)

bikes = []
for b in live['bikes']:
    m, late = merged_detail.get(b['number'], (0, 0))
    bikes.append({**b, 'lapsRule': rule.get(b['number'], 0) + b['adjustment'], 'merged': m, 'late': late})

def ranking(key):
    res = {}
    for g in groups:
        by_team = {}
        for b in bikes:
            if b['group'] == g: by_team.setdefault(b['teamId'], []).append(b)
        rows = []
        for tid, tb in by_team.items():
            tb = sorted(tb, key=lambda b: (-b[key], b['number']))
            best2 = tb[:2]
            rows.append({'team': teams[tid], 'bikes': tb, 'total': sum(b[key] for b in tb),
                         'avg': sum(b[key] for b in best2) / len(best2), 'best': tb[0][key]})
        rows.sort(key=lambda r: (-r['avg'], -r['best'], -r['total']))
        for i, r in enumerate(rows, 1): r['rank'] = i
        res[g] = rows
    return res

app, regle = ranking('laps'), ranking('lapsRule')

# --- Console ------------------------------------------------------------------
now = datetime.now(TZ)
print(f"\nRÉSULTATS — {now:%H:%M:%S}  ·  « Règle » = écart mini {MIN_GAP} s + arrêt à {FIN[11:]}")
for g, name in groups.items():
    print(f"\n{'=' * 100}\n {name.upper()}\n{'=' * 100}")
    print(f" {'Rg règle':>8} {'(app)':>6}  {'Écurie':<24} {'Vélos — règle (app)':<44} {'Moy.2 règle':>11} {'(app)':>7} {'Pts':>5}")
    app_rank = {r['team']['id']: r for r in app[g]}
    for r in regle[g]:
        a = app_rank[r['team']['id']]
        detail = '  '.join(f"#{b['number']}:{b['lapsRule']}" + (f"({b['laps']})" if b['laps'] != b['lapsRule'] else '') for b in r['bikes'])
        flag = '  ⚠️ rang change' if a['rank'] != r['rank'] else ''
        print(f" {r['rank']:>8} {'(' + str(a['rank']) + ')':>6}  {r['team']['name'][:24]:<24} {detail:<44} {r['avg']:>11.1f} {'(' + format(a['avg'], '.1f') + ')':>7} {r['team']['points']:>5}{flag}")

# --- Classeur -----------------------------------------------------------------
sheets = {}
sheets['Résultats'] = [
    {'Classement': groups[g], 'Rang (règle)': r['rank'], 'Rang (app)': next(x['rank'] for x in app[g] if x['team']['id'] == r['team']['id']),
     'Écurie': r['team']['name'], 'Unité': r['team']['unitName'], 'Section': r['team']['sectionName'],
     **{f'Vélo {k + 1}': b['number'] for k, b in enumerate(r['bikes'][:3])},
     **{f'Tours vélo {k + 1} (règle)': b['lapsRule'] for k, b in enumerate(r['bikes'][:3])},
     **{f'Tours vélo {k + 1} (app)': b['laps'] for k, b in enumerate(r['bikes'][:3])},
     'Total écurie (règle)': r['total'], 'Moyenne 2 meilleurs (règle)': round(r['avg'], 2),
     'Total écurie (app)': next(x['total'] for x in app[g] if x['team']['id'] == r['team']['id']),
     'Moyenne 2 meilleurs (app)': round(next(x['avg'] for x in app[g] if x['team']['id'] == r['team']['id']), 2),
     'Points restants': r['team']['points'], 'Points gagnés': r['team']['earned'], 'Points dépensés': r['team']['spent']}
    for g in groups for r in regle[g]]
sheets['Vélos'] = [
    {'Dossard': b['number'], 'Écurie': b['teamName'], 'Classement': groups[b['group']], 'Rang (app)': b['rank'],
     'Tours chrono (app)': b['rawLaps'], 'Bonus/malus/corrections': b['adjustment'], 'Tours (app)': b['laps'],
     f'Relectures < {MIN_GAP} s retirées': b['merged'], 'Passages après la fin retirés': b['late'], 'Tours (règle)': b['lapsRule'],
     'Dernier passage': local(b['lastLapAt'].replace('Z', '') if b['lastLapAt'] else '')}
    for b in sorted(bikes, key=lambda b: b['number'])]
sheets['Écuries'] = [
    {'Écurie': t['name'], 'Unité': t['unitName'], 'Section': t['sectionName'], 'Vélos': ', '.join(map(str, t['bikes'])),
     'Total tours (app)': t['totalLaps'], 'Points gagnés': t['earned'], 'Points dépensés': t['spent'], 'Points restants': t['points']}
    for t in live['teams']]
sheets['Historique points'] = [
    {'Heure': local(r['createdAt']), 'Écurie': r['team'], 'Points': int(r['points']), 'Motif': r['reason'], 'Par': r['performedBy'],
     'Annulé': 'oui' if r['cancelledAt'] else ''}
    for r in sql('SELECT p."createdAt", COALESCE(NULLIF(t."foulardName", \'\'), t."unitName") AS team, p.points, p.reason, p."performedBy", p."cancelledAt" FROM "PointsTransaction" p JOIN "Team" t ON t.id = p."teamId" ORDER BY p."createdAt"')]
sheets['Achats bonus-malus'] = [
    {'Heure': local(r['createdAt']), 'Acheteur': r['buyer'], 'Item': r['itemName'], 'Vélo visé': r['number'], 'Effet (tours)': int(r['lapDelta']),
     'Coût (pts)': int(r['costPoints']), 'Par': r['performedBy'], 'Annulé': 'oui' if r['cancelledAt'] else ''}
    for r in sql('SELECT p."createdAt", COALESCE(NULLIF(t."foulardName", \'\'), t."unitName") AS buyer, p."itemName", d.number, p."lapDelta", p."costPoints", p."performedBy", p."cancelledAt" FROM "Purchase" p JOIN "Team" t ON t.id = p."buyingTeamId" JOIN "Dossard" d ON d.id = p."targetDossardId" ORDER BY p."createdAt"')]
sheets['Corrections'] = [
    {'Heure': local(r['createdAt']), 'Dossard': r['number'], 'Tours': int(r['lapDelta']), 'Motif': r['reason'], 'Par': r['performedBy']}
    for r in sql('SELECT a."createdAt", d.number, a."lapDelta", a.reason, a."performedBy" FROM "RaceAdjustment" a JOIN "Dossard" d ON d.id = a."dossardId" WHERE a.source = \'correction\' ORDER BY a."createdAt"')]
sheets['Passages'] = [
    {'Dossard': int(p['dossardNumber']), 'Heure passage': local(p['timestamp']), 'Reçu à': local(p['createdAt']), 'Tapis': p['mats'], 'Source': p['source'],
     'Après la fin': 'oui' if datetime.fromisoformat(p['timestamp'][:23]) > fin_utc else ''}
    for p in passings]

stamp = now.strftime('%Y%m%d-%H%M')
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        cols = list(dict.fromkeys(k for r in rows for k in r)) or ['(vide)']
        ws.append(cols)
        for r in rows: ws.append([r.get(c, '') for c in cols])
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='E10600')
        for i, c in enumerate(cols, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(10, min(32, len(c) + 2))
        ws.freeze_panes = 'A2'
    path = f'/root/exports/ferme-trophy-export-{stamp}.xlsx'
    wb.save(path); paths = [path]
except ImportError:
    paths = []
    for name, rows in sheets.items():
        path = f"/root/exports/ferme-trophy-{stamp}-{name.replace(' ', '-').lower()}.csv"
        cols = list(dict.fromkeys(k for r in rows for k in r))
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.DictWriter(f, fieldnames=cols or ['vide'], delimiter=';'); w.writeheader(); w.writerows(rows)
        paths.append(path)
print('\nExport :'); [print(f'  scp root@76.13.62.105:{p} ~/Downloads/') for p in paths]
print('  ' + ' · '.join(f'{k}: {len(v)} lignes' for k, v in sheets.items()))
